"""Functional tests: ingest, retrieval, answers, memory, API contract."""
import pytest

import config
import ingest
import rag
from conftest import HEADERS, SID, build_index


# ------------------------------------------------------------------ ingest
def test_chunk_text_respects_size_and_overlap():
    text = "\n".join(f"Line number {i} of the leave policy document." for i in range(200))
    chunks = ingest.chunk_text(text, size=300, overlap=60)
    assert len(chunks) > 1
    assert all(len(c) <= 300 + 60 for c in chunks)
    tail = chunks[0][-40:].split(" ", 1)[1]
    assert tail.split()[0] in chunks[1]  # next chunk starts with the previous tail


def test_chunk_text_splits_one_very_long_line():
    chunks = ingest.chunk_text("word " * 1000, size=200, overlap=0)
    assert len(chunks) > 1 and all(len(c) <= 200 for c in chunks)


def test_chunk_text_empty():
    assert ingest.chunk_text("   \n\n  ") == []


def test_find_files_skips_temp_and_hidden_files():
    names = {p.name for p in ingest.find_files()}
    assert names == {"Leave Policy.md", "Asset Policy.txt", "Holidays.csv"}


def test_csv_rows_become_column_value_lines():
    text = next(ingest.load_csv(config.DOCS_DIR / "Holidays.csv"))[0]
    assert "Date: 2026-01-26 | Holiday: Republic Day" in text


def test_docx_and_xlsx_loaders(tmp_path):
    import pandas as pd
    from docx import Document

    doc = Document()
    doc.add_paragraph("Travel policy paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Class", "Economy"
    doc.save(tmp_path / "t.docx")
    text = next(ingest.load_docx(tmp_path / "t.docx"))[0]
    assert "Travel policy paragraph" in text and "Class | Economy" in text

    pd.DataFrame({"Grade": ["L1"], "Limit": ["500"]}).to_excel(tmp_path / "t.xlsx", sheet_name="Claims", index=False)
    text, location = next(ingest.load_excel(tmp_path / "t.xlsx"))
    assert location == "sheet Claims" and "Grade: L1 | Limit: 500" in text


def test_empty_index_gives_setup_error(llm):
    build_index({})
    try:
        with pytest.raises(RuntimeError, match="No documents are indexed"):
            rag.ask(SID, "leave days")
    finally:
        build_index()


# --------------------------------------------------------------- retrieval
def test_retrieve_finds_the_right_document():
    hits = rag.retrieve("how many days of paid annual leave")
    assert hits and hits[0]["source"] == "Leave Policy.md"


def test_unrelated_question_is_not_found_without_calling_llm(llm):
    result = rag.ask(SID, "zebra quantum xylophone")
    assert result == {
        "answer": "I could not find that in the company documents. Please contact HR or the relevant team.",
        "sources": [],
    }
    assert llm.calls == []


# ----------------------------------------------------------------- answers
def test_answer_with_sources(llm):
    result = rag.ask(SID, "how many days of paid annual leave do employees get")
    assert result["answer"] == "You get 24 days of paid annual leave."
    assert {"source": "Leave Policy.md", "location": ""} in result["sources"]
    assert len(result["sources"]) <= 4
    assert len({(s["source"], s["location"]) for s in result["sources"]}) == len(result["sources"])


def test_prompt_contains_rules_excerpts_and_question(llm):
    rag.ask(SID, "how many days of paid annual leave")
    messages = llm.calls[-1]
    assert messages[0]["role"] == "system" and "Answer ONLY from the policy excerpts" in messages[0]["content"]
    assert "Acme" in messages[0]["content"]
    assert messages[-1]["role"] == "user"
    assert "Policy excerpts:" in messages[-1]["content"] and "(Leave Policy.md)" in messages[-1]["content"]


def test_not_found_marker_is_stripped_and_sources_dropped(llm):
    llm.reply = f"{rag.NOT_FOUND} I could not find that in the company documents."
    result = rag.ask(SID, "annual leave for interns")
    assert result == {"answer": "I could not find that in the company documents.", "sources": []}


def test_sources_capped_at_four(llm, monkeypatch):
    hits = [{"text": "t", "source": f"Doc{i}.pdf", "location": "page 1"} for i in range(7)]
    monkeypatch.setattr(rag, "retrieve", lambda q: hits)
    assert len(rag.ask(SID, "anything")["sources"]) == 4


# ------------------------------------------------------------------ memory
def test_follow_up_is_rewritten_using_history(llm):
    rag.ask(SID, "how many days of annual leave")
    llm.reply = "how many days of annual leave can be carried forward"
    rag.ask(SID, "and carry forward?")
    rewrite_prompt = llm.calls[1][0]["content"]
    assert "Latest question: and carry forward?" in rewrite_prompt
    assert "how many days of annual leave" in rewrite_prompt


def test_first_question_is_not_rewritten(llm):
    rag.ask(SID, "how many days of annual leave")
    assert len(llm.calls) == 1  # answer only, no rewrite call


def test_history_is_trimmed(llm):
    for i in range(5):
        rag.ask(SID, f"annual leave question {i}")
    assert len(rag._history[SID]) == config.HISTORY_TURNS * 2


def test_reset_clears_history(llm):
    rag.ask(SID, "annual leave")
    rag.reset(SID)
    assert SID not in rag._history


def test_sessions_are_isolated(llm):
    rag.ask("b" * 32, "secret question about annual leave")
    rag.ask("c" * 32, "annual leave")
    assert all("secret" not in m["content"] for m in rag._history["c" * 32])
    assert all("secret" not in str(m) for call in llm.calls[1:] for m in call)


def test_session_count_is_bounded(llm, monkeypatch):
    monkeypatch.setattr(rag, "MAX_SESSIONS", 3)
    for i in range(5):
        rag._session(f"s{i}")
    assert len(rag._history) == 3 and "s0" not in rag._history


# --------------------------------------------------------------------- API
def test_health(client):
    assert client.get("/health").get_json() == {"ok": True}


def test_chat_endpoint_contract(client, llm):
    res = client.post("/api/chat", json={"message": "  annual leave days  "}, headers=HEADERS)
    assert res.status_code == 200
    body = res.get_json()
    assert set(body) == {"answer", "sources"}
    assert all(set(s) == {"source", "location"} for s in body["sources"])


@pytest.mark.parametrize(
    "body,message",
    [
        ({"message": ""}, "Type a question first."),
        ({"message": "   "}, "Type a question first."),
        ({}, "Type a question first."),
        ({"message": "x" * 1001}, "Keep your question under 1,000 characters."),
    ],
)
def test_chat_validation(client, llm, body, message):
    res = client.post("/api/chat", json=body, headers=HEADERS)
    assert res.status_code == 400 and res.get_json() == {"error": message}


def test_exactly_1000_characters_is_allowed(client, llm):
    assert client.post("/api/chat", json={"message": "leave " + "x" * 994}, headers=HEADERS).status_code == 200


def test_setup_error_is_shown(client, monkeypatch):
    def boom(sid, q):
        raise RuntimeError("XAI_API_KEY is missing. Copy .env.example to .env and add your xAI key.")

    monkeypatch.setattr(rag, "ask", boom)
    res = client.post("/api/chat", json={"message": "hi"}, headers=HEADERS)
    assert res.status_code == 500 and "XAI_API_KEY is missing" in res.get_json()["error"]


def test_zscaler_error_is_explained(client, monkeypatch):
    def boom(sid, q):
        raise Exception("Blocked by Zscaler")

    monkeypatch.setattr(rag, "ask", boom)
    res = client.post("/api/chat", json={"message": "hi"}, headers=HEADERS)
    assert "company network proxy (Zscaler)" in res.get_json()["error"]


def test_other_errors_show_type_and_short_message(client, monkeypatch):
    def boom(sid, q):
        raise ValueError("model not found " + "x" * 1000)

    monkeypatch.setattr(rag, "ask", boom)
    error = client.post("/api/chat", json={"message": "hi"}, headers=HEADERS).get_json()["error"]
    assert error.startswith("Lookup failed: ValueError: model not found") and len(error) < 360


def test_reset_endpoint(client, llm):
    client.post("/api/chat", json={"message": "annual leave"}, headers=HEADERS)
    assert client.post("/api/reset", headers=HEADERS).get_json() == {"ok": True}
    assert SID not in rag._history


# ------------------------------------------------------ review fixes
def test_history_turns_zero_turns_memory_off(llm, monkeypatch):
    monkeypatch.setattr(config, "HISTORY_TURNS", 0)
    for i in range(3):
        rag.ask(SID, f"annual leave question {i}")
    assert rag._history[SID] == []
    assert all(len(call) == 2 for call in llm.calls)  # system + question only, no rewrite calls


def test_active_session_is_not_evicted_lru(llm, monkeypatch):
    monkeypatch.setattr(rag, "MAX_SESSIONS", 3)
    rag._session("old-but-active")
    rag._session("s1")
    rag._session("s2")
    rag._session("old-but-active")  # used again: becomes most recent
    rag._session("s3")
    assert "old-but-active" in rag._history and "s1" not in rag._history


def test_empty_model_reply_becomes_not_found(llm):
    llm.reply = "   "
    assert rag.ask(SID, "annual leave") == {"answer": rag.NOT_FOUND_ANSWER, "sources": []}


def test_invalid_embed_provider_fails_fast(monkeypatch):
    import importlib

    monkeypatch.setenv("EMBED_PROVIDER", "OpenAI ")
    assert importlib.reload(config).EMBED_PROVIDER == "openai"
    monkeypatch.setenv("EMBED_PROVIDER", "bogus")
    with pytest.raises(RuntimeError, match="EMBED_PROVIDER must be"):
        importlib.reload(config)
    monkeypatch.setenv("EMBED_PROVIDER", "keyword")
    importlib.reload(config)


def test_config_endpoint(client):
    assert client.get("/api/config", headers=HEADERS).get_json() == {"orgName": "Acme"}
    assert client.get("/api/config").status_code == 403


def test_llm_client_has_timeout_and_reply_cap(monkeypatch):
    import llm

    seen = {}

    class FakeCompletions:
        def create(self, **kw):
            seen.update(kw)
            return type("R", (), {"choices": [type("C", (), {"message": type("M", (), {"content": " hi "})})]})

    monkeypatch.setattr(llm, "_chat_client", None)
    monkeypatch.setattr(config, "XAI_API_KEY", "k")
    client = llm._get_chat_client()
    assert client.timeout == config.LLM_TIMEOUT and client.max_retries == 1
    monkeypatch.setattr(llm, "_chat_client", type("X", (), {"chat": type("Y", (), {"completions": FakeCompletions()})}))
    assert llm.chat([{"role": "user", "content": "q"}]) == "hi"
    assert seen["max_tokens"] == config.MAX_TOKENS
    monkeypatch.setattr(llm, "_chat_client", None)
