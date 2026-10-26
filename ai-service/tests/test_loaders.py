"""Every supported file format is read, chunked and labelled correctly."""
import email.message
import zipfile

import pandas as pd
import pytest

import config
import ingest
from conftest import DOCS, build_index


def texts(path):
    return [(t, loc) for t, loc in ingest.LOADERS[path.suffix.lower()](path)]


def joined(path):
    return "\n".join(t for t, _ in texts(path))


def test_pptx_slides_tables_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Travel policy"
    s1.notes_slide.notes_text_frame.text = "Speaker note: economy class only"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    table = s2.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Grade", "Limit"
    prs.save(tmp_path / "t.pptx")
    out = texts(tmp_path / "t.pptx")
    assert out[0][1] == "slide 1" and "Travel policy" in out[0][0] and "economy class only" in out[0][0]
    assert out[1][1] == "slide 2" and "Grade | Limit" in out[1][0]


def test_odt_and_odp(tmp_path):
    from odf.opendocument import OpenDocumentPresentation, OpenDocumentText
    from odf.text import H, P

    doc = OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text="Leave"))
    doc.text.addElement(P(text="24 days per year"))
    doc.save(str(tmp_path / "t.odt"))
    assert "Leave" in joined(tmp_path / "t.odt") and "24 days per year" in joined(tmp_path / "t.odt")

    from odf.draw import Frame, Page, TextBox
    from odf.style import MasterPage, PageLayout

    pres = OpenDocumentPresentation()
    layout = PageLayout(name="L")
    pres.automaticstyles.addElement(layout)
    master = MasterPage(name="M", pagelayoutname=layout)
    pres.masterstyles.addElement(master)
    page = Page(masterpagename=master)
    frame = Frame(width="10cm", height="2cm", x="1cm", y="1cm")
    box = TextBox()
    box.addElement(P(text="Slide text"))
    frame.addElement(box)
    page.addElement(frame)
    pres.presentation.addElement(page)
    pres.save(str(tmp_path / "t.odp"))
    assert "Slide text" in joined(tmp_path / "t.odp")


def test_ods_xlsx_and_tsv(tmp_path):
    df = pd.DataFrame({"Grade": ["L1"], "Limit": ["500"]})
    df.to_excel(tmp_path / "t.ods", sheet_name="Claims", index=False, engine="odf")
    assert texts(tmp_path / "t.ods") == [("Grade: L1 | Limit: 500", "sheet Claims")]
    (tmp_path / "t.tsv").write_text("Grade\tLimit\nL2\t900\n")
    assert joined(tmp_path / "t.tsv") == "Grade: L2 | Limit: 900"


def test_rtf(tmp_path):
    (tmp_path / "t.rtf").write_text(r"{\rtf1\ansi {\b Notice period} is 30 days.\par}")
    assert "Notice period is 30 days." in joined(tmp_path / "t.rtf")


def test_html_drops_scripts_and_keeps_table_rows(tmp_path):
    (tmp_path / "t.html").write_text(
        "<html><head><style>p{}</style><script>alert('x')</script></head><body>"
        "<h1>Benefits</h1><p>Health cover</p><table><tr><th>Plan</th><td>Gold</td></tr></table></body></html>"
    )
    text = joined(tmp_path / "t.html")
    assert "Benefits" in text and "Health cover" in text and "Plan | Gold" in text
    assert "alert" not in text and "p{}" not in text


def test_xml(tmp_path):
    (tmp_path / "t.xml").write_text("<policy><title>WFH</title><rule>2 days a week</rule></policy>")
    assert joined(tmp_path / "t.xml") == "WFH\n2 days a week"


def test_xml_entity_bomb_is_refused(tmp_path):
    bomb = (tmp_path / "bomb.xml")
    bomb.write_text('<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]><l>&b;</l>')
    with pytest.raises(Exception):
        joined(bomb)


def test_json_keeps_key_paths(tmp_path):
    (tmp_path / "t.json").write_text('{"leave": {"annual": 24, "types": ["sick", "casual"]}, "empty": null}')
    assert joined(tmp_path / "t.json") == "leave.annual: 24\nleave.types[0]: sick\nleave.types[1]: casual"


@pytest.mark.parametrize("ext", [".txt", ".md", ".markdown", ".rst", ".log", ".yaml", ".yml"])
def test_plain_text_formats(tmp_path, ext):
    (tmp_path / f"t{ext}").write_text("Dress code: business casual")
    assert joined(tmp_path / f"t{ext}") == "Dress code: business casual"


def test_epub_sections(tmp_path):
    with zipfile.ZipFile(tmp_path / "t.epub", "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("OEBPS/ch1.xhtml", "<html><body><p>Chapter one: onboarding</p></body></html>")
        z.writestr("OEBPS/ch2.xhtml", "<html><body><p>Chapter two: exit</p></body></html>")
    assert texts(tmp_path / "t.epub") == [("Chapter one: onboarding", "section 1"), ("Chapter two: exit", "section 2")]


def test_epub_zip_bomb_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "EPUB_MAX_BYTES", 100)
    with zipfile.ZipFile(tmp_path / "big.epub", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("a.xhtml", "x" * 1000)
    with pytest.raises(ValueError, match="too large"):
        joined(tmp_path / "big.epub")


def test_eml_plain_and_html(tmp_path):
    msg = email.message.EmailMessage()
    msg["Subject"], msg["From"] = "Holiday list 2026", "hr@acme.test"
    msg.set_content("Offices close on 26 January.")
    (tmp_path / "t.eml").write_bytes(bytes(msg))
    text = joined(tmp_path / "t.eml")
    assert "Subject: Holiday list 2026" in text and "Offices close on 26 January." in text

    html = email.message.EmailMessage()
    html["Subject"] = "HTML mail"
    html.set_content("<p>Bring your <b>badge</b></p><script>x()</script>", subtype="html")
    (tmp_path / "h.eml").write_bytes(bytes(html))
    assert "Bring your\nbadge" in joined(tmp_path / "h.eml") and "x()" not in joined(tmp_path / "h.eml")


def test_every_extension_has_a_loader():
    expected = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xlsm", ".xls", ".ods", ".odt", ".odp",
                ".csv", ".tsv", ".txt", ".md", ".markdown", ".rst", ".log", ".yaml", ".yml", ".rtf", ".html",
                ".htm", ".xml", ".json", ".epub", ".eml", ".msg", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
    assert set(ingest.LOADERS) == expected


# ---------------------------------------------------- ingest-level behaviour
def _ingest(files, capsys):
    """Run ingest on `files` (name -> bytes|str) and return (printed output, indexed sources)."""
    import store

    for f in config.DOCS_DIR.iterdir():
        f.unlink()
    for name, data in files.items():
        p = config.DOCS_DIR / name
        p.write_bytes(data) if isinstance(data, bytes) else p.write_text(data)
    capsys.readouterr()
    ingest.main()
    out = capsys.readouterr().out
    got = store.get_collection().get(include=["metadatas"])["metadatas"]
    return out, {m["source"] for m in got}


@pytest.fixture
def restore_index():
    yield
    build_index()


def test_image_and_scanned_pdf_without_ocr_are_reported(capsys, restore_index, monkeypatch):
    from io import BytesIO

    from pypdf import PdfWriter

    monkeypatch.setattr(ingest, "_ocr_available", lambda: False)
    buf = BytesIO()
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.write(buf)
    out, sources = _ingest({"scan.pdf": buf.getvalue(), "photo.png": b"\x89PNG", "ok.txt": "annual leave"}, capsys)
    assert "scan.pdf: no text found" in out
    assert "Skipped photo.png: images need OCR" in out
    assert sources == {"ok.txt"}


def test_legacy_doc_without_libreoffice_is_reported(capsys, restore_index, monkeypatch):
    monkeypatch.setattr(ingest.shutil, "which", lambda name: None)
    out, sources = _ingest({"old.doc": b"\xd0\xcf\x11\xe0", "ok.txt": "annual leave"}, capsys)
    assert "Skipped old.doc: install LibreOffice" in out and sources == {"ok.txt"}


def test_symlinks_are_not_followed(capsys, restore_index, tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("salary of the CEO")
    out, sources = _ingest({"ok.txt": "annual leave"}, capsys)
    (config.DOCS_DIR / "link.txt").symlink_to(secret)
    ingest.main()
    import store

    assert {m["source"] for m in store.get_collection().get(include=["metadatas"])["metadatas"]} == {"ok.txt"}


def test_files_over_the_size_cap_are_skipped(capsys, restore_index, monkeypatch):
    monkeypatch.setattr(config, "MAX_FILE_MB", 0)
    with pytest.raises(SystemExit):  # the only file was skipped, so the old index is kept
        _ingest({"big.txt": "annual leave"}, capsys)
    assert "Skipped big.txt: larger than 0 MB" in capsys.readouterr().out


def test_unreadable_files_keep_the_old_index(capsys, restore_index):
    import rag

    for f in config.DOCS_DIR.iterdir():
        f.unlink()
    bad = config.DOCS_DIR / "locked.txt"
    bad.write_text("annual leave")
    bad.chmod(0)
    try:
        with pytest.raises(SystemExit, match="existing index was kept"):
            ingest.main()
    finally:
        bad.chmod(0o644)
    assert rag.retrieve("paid annual leave")[0]["source"] == "Leave Policy.md"


def test_stale_pointer_does_not_recreate_deleted_collection(restore_index, monkeypatch):
    import store

    old = store._active_name()
    names = iter([old, None])  # first read is stale, second read sees the new pointer
    build_index()
    new = store._active_name()
    real = store._active_name
    monkeypatch.setattr(store, "_active_name", lambda: next(names) or real())
    assert store.get_collection().name == new
    listed = [c if isinstance(c, str) else c.name for c in store._client().list_collections()]
    assert old not in listed


def test_file_failing_halfway_adds_no_chunks(capsys, restore_index, monkeypatch):
    def half(path):
        yield "first part of the policy", "page 1"
        raise ValueError("corrupt page 2")

    monkeypatch.setitem(ingest.LOADERS, ".txt", half)
    out, sources = _ingest({"broken.txt": "x", "fine.md": "annual leave"}, capsys)
    assert "Skipped broken.txt: corrupt page 2" in out and sources == {"fine.md"}


def test_failed_rebuild_keeps_the_old_index(capsys, restore_index, monkeypatch):
    import rag

    def broken_embed(texts):
        raise RuntimeError("embedding service down")

    monkeypatch.setattr(ingest, "embed", broken_embed)
    with pytest.raises(RuntimeError):
        ingest.main()
    assert rag.retrieve("paid annual leave")[0]["source"] == "Leave Policy.md"


def test_rebuild_removes_old_collections(restore_index):
    import store

    build_index()
    build_index()
    names = [c if isinstance(c, str) else c.name for c in store._client().list_collections()]
    assert len(names) == 1 and names[0] == store._active_name()


# ------------------------------------------------------- round-2 review fixes
@pytest.mark.parametrize(
    "data",
    [
        "Café résumé — price €100".encode("cp1252"),
        "Café résumé — price €100".encode("utf-8"),
        "\ufeffCafé résumé — price €100".encode("utf-8"),
        "Café résumé — price €100".encode("utf-16"),  # Notepad "Unicode" (with BOM)
        "Café résumé — price €100".encode("utf-16-le"),  # no BOM
    ],
)
def test_text_encodings_are_not_mangled(tmp_path, data):
    (tmp_path / "t.txt").write_bytes(data)
    assert joined(tmp_path / "t.txt") == "Café résumé — price €100"


def test_html_honours_its_charset(tmp_path):
    (tmp_path / "t.html").write_bytes('<meta charset="windows-1252"><p>Café €5</p>'.encode("cp1252"))
    assert joined(tmp_path / "t.html") == "Café €5"


def test_odt_keeps_headings_in_order(tmp_path):
    from odf.opendocument import OpenDocumentText
    from odf.text import H, List, ListItem, P

    doc = OpenDocumentText()
    for title, body in (("Section A", "Body A"), ("Section B", "Body B")):
        doc.text.addElement(H(outlinelevel=1, text=title))
        doc.text.addElement(P(text=body))
    lst, item = List(), ListItem()
    item.addElement(P(text="List item C"))
    lst.addElement(item)
    doc.text.addElement(lst)
    doc.save(str(tmp_path / "o.odt"))
    assert joined(tmp_path / "o.odt") == "Section A\nBody A\nSection B\nBody B\nList item C"


def test_docx_reads_content_controls_nested_tables_and_tracked_inserts(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = Document()
    doc.add_paragraph("First paragraph")
    body = doc.element.body
    sect = body[-1]  # keep sectPr last
    sect.addprevious(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Inside content control</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
    sect.addprevious(parse_xml(
        f'<w:p {nsdecls("w")}><w:ins w:id="1" w:author="a"><w:r><w:t>Tracked insert</w:t></w:r></w:ins>'
        f'<w:del w:id="2" w:author="a"><w:r><w:delText>Deleted text</w:delText></w:r></w:del></w:p>'))
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Outer"
    outer.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested cell"
    doc.add_paragraph("Last paragraph")
    doc.save(tmp_path / "c.docx")
    lines = joined(tmp_path / "c.docx").splitlines()
    assert lines[0] == "First paragraph" and lines[-1] == "Last paragraph"  # document order
    text = "\n".join(lines)
    assert "Inside content control" in text and "Tracked insert" in text and "Deleted text" not in text
    assert any(l.startswith("Outer | ") and "Nested cell" in l for l in lines)


def test_pptx_group_shapes_and_missing_notes_body(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Grouped text"
    notes = slide.notes_slide
    for ph in list(notes.placeholders):
        if ph.placeholder_format.type == 2:  # BODY
            ph._element.getparent().remove(ph._element)
    prs.save(tmp_path / "g.pptx")
    assert texts(tmp_path / "g.pptx") == [("Grouped text", "slide 1")]


def test_multi_page_tiff_ocrs_every_page(tmp_path, monkeypatch):
    from PIL import Image

    pages = [Image.new("L", (10, 10), c) for c in (0, 128, 255)]
    pages[0].save(tmp_path / "s.tiff", save_all=True, append_images=pages[1:])
    monkeypatch.setattr(ingest, "_ocr_available", lambda: True)
    monkeypatch.setattr(ingest, "_ocr_image", lambda img: f"colour {img.getpixel((0, 0))}")
    assert texts(tmp_path / "s.tiff") == [("colour 0", "page 1"), ("colour 128", "page 2"), ("colour 255", "page 3")]


def test_eml_attachments_are_indexed(tmp_path, capsys):
    msg = email.message.EmailMessage()
    msg["Subject"] = "Policy update"
    msg.set_content("See attached.")
    msg.add_attachment(b"Carry forward limit is 5 days.", maintype="text", subtype="plain", filename="../../evil name.txt")
    msg.add_attachment(b"\x00\x01", maintype="application", subtype="octet-stream", filename="tool.exe")
    (tmp_path / "a.eml").write_bytes(bytes(msg))
    out = texts(tmp_path / "a.eml")
    assert ("Carry forward limit is 5 days.", "attachment evil name.txt") in out
    assert "tool.exe: not indexed" in capsys.readouterr().out


def test_libreoffice_timeout_kills_the_whole_process_group(tmp_path, monkeypatch):
    import subprocess

    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\nsleep 300 &\necho $! > \"$PIDFILE\"\nwait\n")
    fake.chmod(0o755)
    pidfile = tmp_path / "child.pid"
    monkeypatch.setenv("PIDFILE", str(pidfile))
    monkeypatch.setattr(ingest.shutil, "which", lambda name: str(fake) if name == "soffice" else None)
    real_popen = subprocess.Popen

    class QuickPopen(real_popen):
        def wait(self, timeout=None):
            return super().wait(timeout=1 if timeout else None)

    monkeypatch.setattr(ingest.subprocess, "Popen", QuickPopen)
    (tmp_path / "x.doc").write_bytes(b"x")
    with pytest.raises(RuntimeError, match="longer than 120 seconds"):
        joined(tmp_path / "x.doc")
    import os
    import time

    child = int(pidfile.read_text())
    time.sleep(0.2)
    with pytest.raises(ProcessLookupError):
        os.kill(child, 0)
