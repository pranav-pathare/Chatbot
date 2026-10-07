"""Read every file in docs/, cut it into chunks, embed the chunks, save them to Chroma.

Run:   python ingest.py
Re-run whenever you add or change documents. It rebuilds the index from scratch.
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader

import config
from llm import embed
from store import reset_collection


# ---------------------------------------------------------------- loaders
# Each loader yields (text, location) pairs. `location` is shown to the
# employee as part of the source tag, e.g. "Leave Policy.pdf, page 3".

def load_pdf(path: Path):
    reader = PdfReader(str(path))
    for number, page in enumerate(reader.pages, start=1):
        yield page.extract_text() or "", f"page {number}"


def load_docx(path: Path):
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    yield "\n".join(parts), ""


def _frame_to_text(df: pd.DataFrame) -> str:
    """One line per row, written as 'Column: value | Column: value'."""
    df = df.dropna(how="all").fillna("")
    lines = []
    for _, row in df.iterrows():
        pairs = [f"{col}: {str(val).strip()}" for col, val in row.items() if str(val).strip()]
        if pairs:
            lines.append(" | ".join(pairs))
    return "\n".join(lines)


def load_excel(path: Path):
    for sheet_name, df in pd.read_excel(path, sheet_name=None, dtype=str).items():
        yield _frame_to_text(df), f"sheet {sheet_name}"


def load_csv(path: Path):
    yield _frame_to_text(pd.read_csv(path, dtype=str)), ""


def load_text(path: Path):
    yield path.read_text(encoding="utf-8", errors="ignore"), ""


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".xlsx": load_excel,
    ".xls": load_excel,
    ".csv": load_csv,
    ".txt": load_text,
    ".md": load_text,
}


# --------------------------------------------------------------- chunking
def chunk_text(text: str, size: int = config.CHUNK_SIZE, overlap: int = config.CHUNK_OVERLAP):
    """Pack whole lines into chunks of roughly `size` characters.

    Lines are never cut mid-way unless one line alone is longer than `size`.
    Each new chunk starts with the tail of the previous one (`overlap`) so a
    sentence on a boundary is not lost.
    """
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    pieces = []
    for line in filter(None, lines):
        while len(line) > size:  # a very long line: split at a sentence or space
            cut = line.rfind(". ", 0, size)
            if cut == -1:
                cut = line.rfind(" ", 0, size)
            if cut == -1:
                cut = size - 1
            pieces.append(line[: cut + 1].strip())
            line = line[cut + 1 :].strip()
        if line:
            pieces.append(line)

    chunks, current = [], ""
    for piece in pieces:
        if current and len(current) + len(piece) + 1 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            tail = tail[tail.find(" ") + 1 :] if " " in tail else ""  # start on a word
            current = f"{tail} {piece}".strip()
        else:
            current = f"{current}\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks


# ------------------------------------------------------------------- main
def find_files():
    return sorted(
        p
        for p in config.DOCS_DIR.rglob("*")
        if p.is_file()
        and p.suffix.lower() in LOADERS
        and not p.name.startswith(("~$", "."))  # skip Office temp files and hidden files
    )


def collect_records():
    records = []
    for path in find_files():
        name = str(path.relative_to(config.DOCS_DIR))
        count = 0
        try:
            for text, location in LOADERS[path.suffix.lower()](path):
                for chunk in chunk_text(text):
                    records.append({"text": chunk, "source": name, "location": location})
                    count += 1
        except Exception as error:
            print(f"  ! Skipped {name}: {error}")
            continue
        if count == 0:
            print(f"  ! {name}: no text found (a scanned PDF needs OCR first)")
        else:
            print(f"  + {name}: {count} chunks")
    return records


def main():
    config.DOCS_DIR.mkdir(exist_ok=True)
    print(f"Reading documents from {config.DOCS_DIR}")
    records = collect_records()
    if not records:
        print("Nothing to index. Add PDF, Word, Excel, CSV or text files to docs/ and run again.")
        return

    print(f"Embedding {len(records)} chunks...")
    # The file name goes into the embedded text so a question like
    # "what does the asset policy say" can match on the title too.
    titled = [f"{Path(r['source']).stem}\n{r['text']}" for r in records]
    vectors = embed(titled)

    collection = reset_collection()
    step = 500
    for i in range(0, len(records), step):
        part = records[i : i + step]
        collection.add(
            ids=[f"chunk-{i + j}" for j in range(len(part))],
            documents=[r["text"] for r in part],
            embeddings=vectors[i : i + step],
            metadatas=[{"source": r["source"], "location": r["location"]} for r in part],
        )
    print(f"Done. {collection.count()} chunks indexed.")


if __name__ == "__main__":
    main()
